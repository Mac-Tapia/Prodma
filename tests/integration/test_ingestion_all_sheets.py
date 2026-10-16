from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from prodma.artifacts import RunLayout
from prodma.config import PipelineConfig
from prodma.custody.inventory import CustodyBlockedError
from prodma.stages.ingestion import run_ingestion


def _demand_workbook(path: Path, sheets: tuple[str, ...]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name in sheets:
        sheet = workbook.create_sheet(name)
        sheet.append(["Reporte de medidor"])
        sheet.append(["Local Time", "kW tot mean", "kVAR tot mean", "kVA tot mean"])
        sheet.append([pd.Timestamp("2030-01-01 00:15"), 100.0, 25.0, 103.08])
    workbook.save(path)
    return path


def _catalog(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "alimentador_medidor"
    sheet.append(["codigo alimentador", "Medidor", "Nivel Tension", "siglas"])
    sheet.append(["LT", "MED-01", 60, None])
    workbook.save(path)
    return path


def _config(root: Path, inputs: dict) -> PipelineConfig:
    return PipelineConfig(
        project_root=root,
        output_root=root / "resultados",
        frequency="15min",
        timezone="America/Lima",
        seed=42,
        horizons=(96,),
        raw={"inputs": {"historical_dir": "datahistorica", **inputs}},
    )


def test_future_year_and_every_sheet_are_discovered_without_configured_years(tmp_path):
    _demand_workbook(tmp_path / "datahistorica" / "demanda2030.xlsx", sheets=("MED-01", "MED-02"))
    layout = RunLayout.create(tmp_path / "resultados", "r-futuro")

    result = run_ingestion(_config(tmp_path, {}), layout)

    assert len(result.artifacts["parquet"]) == 2
    assert result.status == "SUCCESS"


def test_publishes_custody_inventory_and_authorized_catalog(tmp_path):
    _demand_workbook(tmp_path / "datahistorica" / "demanda2030.xlsx", sheets=("MED-01",))
    _catalog(tmp_path / "medidoralimentador.xlsx")
    layout = RunLayout.create(tmp_path / "resultados", "r-custodia")

    result = run_ingestion(_config(tmp_path, {"authorized_catalog": "medidoralimentador.xlsx"}), layout)

    custody = layout.root / "00_custodia"
    assert (custody / "inventario.jsonl").exists()
    assert (custody / "input_summary.json").exists()
    reopened = pd.read_parquet(custody / "authorized_meters.parquet")
    assert list(reopened["identificador"]) == ["MED-01"]
    assert result.artifacts["custodia"]["catalogo"].endswith("authorized_meters.parquet")


def test_missing_mandatory_catalog_fails_closed_before_ingestion(tmp_path):
    _demand_workbook(tmp_path / "datahistorica" / "demanda2030.xlsx", sheets=("MED-01",))
    layout = RunLayout.create(tmp_path / "resultados", "r-bloqueo")

    with pytest.raises(CustodyBlockedError, match="medidoralimentador.xlsx"):
        run_ingestion(_config(tmp_path, {"authorized_catalog": "medidoralimentador.xlsx"}), layout)

    assert list(layout.stage("01_ingesta", "parquet").glob("*.parquet")) == []
