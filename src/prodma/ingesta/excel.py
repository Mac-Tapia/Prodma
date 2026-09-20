"""Capa de ingesta ("bronce"): libros Excel de medición -> un parquet estandarizado por hoja.

Misma semántica que la sección 8 del notebook (encabezado en las primeras 20 filas, primera
columna ante encabezados repetidos, fechas alineadas a 15 min con tolerancia, duplicados
resueltos sin promediar contradicciones), pero conservando además las magnitudes de estado del
medidor —tensión, corriente, frecuencia, desbalance— que permiten distinguir "el equipo no tenía
energía / no medía" de "la carga era realmente cero".

Es la única capa que toca los Excel. Todo lo demás (calidad, dataset modelable) trabaja sobre estos
parquet, sin volver a leer los libros.
"""
from __future__ import annotations

import gc
import hashlib
import json
import re
from pathlib import Path
from typing import Any, Iterator

import numpy as np
import pandas as pd

from prodma.excel_schema import find_first_columns, normalize_header
from prodma.series import resolver_duplicados

VERSION_PARSER = "bronce_v1"
FRECUENCIA = "15min"
ALIAS_MEDIDAS: dict[str, set[str]] = {
    "P": {"kw tot", "kw tot mean"},
    "Q": {"kvar tot", "kvar tot mean"},
    "S_medida": {"kva tot", "kva tot mean", "kva"},
    "V_prom": {"vll avg", "vll avg mean"},
    "I_prom": {"i avg", "i avg mean"},
    "Frecuencia": {"freq", "freq mean"},
    "V_desbalance": {"v unbal", "v unbal mean"},
    "Fecha": {"local time"},
    "UTC": {"timestamp"},
}
COLUMNAS_MEDIDA = ("P", "Q", "S_medida", "V_prom", "I_prom", "Frecuencia", "V_desbalance")


def abrir_libro(ruta) -> tuple[str, Any]:
    """Abre .xls (xlrd) o .xlsx (openpyxl) según la extensión real del archivo."""
    ext = Path(ruta).suffix.lower()
    if ext == ".xls":
        import xlrd

        return "xls", xlrd.open_workbook(str(ruta), on_demand=True)
    if ext == ".xlsx":
        import openpyxl

        return "xlsx", openpyxl.load_workbook(str(ruta), read_only=True, data_only=True)
    raise ValueError(f"Extensión no soportada: {ext} (solo .xls o .xlsx)")


def nombres_hojas(tipo: str, wb) -> list[str]:
    return list(wb.sheet_names()) if tipo == "xls" else list(wb.sheetnames)


def filas_hoja(tipo: str, wb, nombre: str) -> Iterator[list]:
    if tipo == "xls":
        hoja = wb.sheet_by_name(nombre)
        for r in range(hoja.nrows):
            yield hoja.row_values(r)
    else:
        for fila in wb[nombre].iter_rows(values_only=True):
            yield list(fila)


def cerrar_libro(tipo: str, wb) -> None:
    if tipo == "xls":
        wb.release_resources()
    else:
        wb.close()


def _liberar_hoja(tipo: str, wb, nombre: str) -> None:
    if tipo == "xls":
        wb.unload_sheet(nombre)


def fechas_excel(valores, datemode: int = 0) -> pd.DatetimeIndex:
    def una(v):
        if v is None or v == "":
            return pd.NaT
        try:
            if isinstance(v, (float, int, np.number)) and not isinstance(v, bool):
                import xlrd

                return xlrd.xldate_as_datetime(float(v), datemode)
            return pd.to_datetime(v, dayfirst=True, errors="coerce")
        except (ValueError, OverflowError):
            return pd.NaT

    return pd.DatetimeIndex([una(v) for v in valores])


def alinear_fechas(idx: pd.DatetimeIndex, tolerancia_seg: float = 10):
    """(fecha redondeada a 15 min, válida, desvío en segundos). Válida = existe y está a <= tolerancia."""
    redondeado = idx.round(FRECUENCIA)
    delta = np.abs((idx - redondeado).total_seconds())
    valida = idx.notna() & (delta <= tolerancia_seg)
    return redondeado, valida, delta


def leer_hoja(tipo: str, wb, nombre: str, tolerancia_seg: float = 10) -> tuple[pd.DataFrame, dict]:
    """Una hoja -> (DataFrame indexado por Fecha con las columnas de medida disponibles, info de auditoría)."""
    encabezado, cabecera, datos = None, None, []
    for i, fila in enumerate(filas_hoja(tipo, wb, nombre)):
        if encabezado is None:
            if i >= 20:
                break
            if "local time" in [normalize_header(x) for x in fila]:
                encabezado, cabecera = i, fila
            continue
        datos.append(fila)
    if encabezado is None:
        raise ValueError("Encabezado no encontrado")
    cols = find_first_columns(cabecera, ALIAS_MEDIDAS, required={"Fecha", "P", "Q"})
    crudo = {k: [(f[c] if c < len(f) else None) for f in datos] for k, c in cols.items()}
    datemode = getattr(wb, "datemode", 0) if tipo == "xls" else 0
    original = fechas_excel(crudo.pop("Fecha"), datemode)
    dt, valida, delta = alinear_fechas(original, tolerancia_seg)
    utc = fechas_excel(crudo.pop("UTC"), datemode) if "UTC" in crudo else None
    df = pd.DataFrame({k: pd.to_numeric(pd.Series(v), errors="coerce").to_numpy() for k, v in crudo.items()}, index=dt)
    df.index.name = "Fecha"
    df = df.loc[valida].sort_index()
    repetidas = int(df.index.duplicated(keep=False).sum())
    df = resolver_duplicados(df)
    info = dict(
        hoja=nombre, filas=len(df), inicio=str(df.index.min()) if len(df) else None,
        fin=str(df.index.max()) if len(df) else None, columnas_medida=[c for c in COLUMNAS_MEDIDA if c in df.columns],
        fechas_desalineadas=int((~valida).sum()), fechas_redondeadas=int(((delta > 0) & valida).sum()),
        filas_duplicadas=repetidas,
        UTC_desfase_h=float(np.nanmedian((utc - original).total_seconds() / 3600)) if utc is not None else None,
        intervalos_ausentes=(len(pd.date_range(df.index.min(), df.index.max(), freq=FRECUENCIA)) - len(df)) if len(df) else 0,
    )
    return df.astype("float64"), info


def huella_libro(ruta) -> str:
    st = Path(ruta).stat()
    return f"{Path(ruta).name}|{st.st_size}|{st.st_mtime_ns}"


def ingestar_libro(ruta, destino, tolerancia_seg: float = 10, imprimir=print,
                   destino_auditoria=None) -> list[dict]:
    """Ingesta un libro anual: un parquet + un json por hoja de medidor (caché por huella del libro).

    Devuelve un registro por hoja: ``archivo, anio, medidor, estado, path, info`` (o ``error``).
    """
    ruta, destino = Path(ruta), Path(destino)
    destino_auditoria = Path(destino_auditoria) if destino_auditoria is not None else destino
    destino.mkdir(parents=True, exist_ok=True)
    destino_auditoria.mkdir(parents=True, exist_ok=True)
    m = re.fullmatch(r"demanda(\d{4})\.xlsx?", ruta.name)
    anio = int(m.group(1)) if m else None
    huella = huella_libro(ruta)
    registros = []
    tipo, wb = abrir_libro(ruta)
    try:
        for hoja in nombres_hojas(tipo, wb):
            if hoja.startswith("|PRIVATE|"):
                continue
            ck = hashlib.sha256(f"{huella}|{hoja}|{VERSION_PARSER}|{tolerancia_seg}".encode()).hexdigest()[:24]
            pq, js = destino / f"{ck}.parquet", destino_auditoria / f"{ck}.json"
            reg = dict(archivo=ruta.name, anio=anio, medidor=hoja, path=str(pq))
            try:
                if pq.exists() and js.exists():
                    info = json.loads(js.read_text(encoding="utf-8"))
                else:
                    df, info = leer_hoja(tipo, wb, hoja, tolerancia_seg)
                    info.update(archivo=ruta.name, anio=anio)
                    df.to_parquet(pq)
                    js.write_text(json.dumps(info, ensure_ascii=False, default=float), encoding="utf-8")
                reg.update(estado="ok", info=info)
            except Exception as exc:  # una hoja mala no debe tumbar el libro: se registra
                reg.update(estado="error", error=f"{type(exc).__name__}: {exc}", info={})
            registros.append(reg)
            _liberar_hoja(tipo, wb, hoja)
            gc.collect()
    finally:
        cerrar_libro(tipo, wb)
    imprimir(f"{ruta.name}: {sum(r['estado'] == 'ok' for r in registros)} hojas ok, "
             f"{sum(r['estado'] == 'error' for r in registros)} con error")
    return registros


def unir_por_medidor(registros: list[dict], inicio: str | None = None) -> dict[str, pd.DataFrame]:
    """Une los años de cada medidor en una serie con grilla completa de 15 min (NaN donde no hay dato).

    Las marcas repetidas entre libros (empalme de años) se resuelven con la misma regla que dentro de
    una hoja. La grilla va de ``inicio`` (o la primera marca del medidor) a su última marca.
    """
    por_medidor: dict[str, list[pd.DataFrame]] = {}
    for r in registros:
        if r["estado"] == "ok":
            por_medidor.setdefault(r["medidor"], []).append(pd.read_parquet(r["path"]))
    salida = {}
    for medidor, partes in por_medidor.items():
        df = resolver_duplicados(pd.concat(partes).sort_index())
        ini = pd.Timestamp(inicio) if inicio else df.index.min()
        salida[medidor] = df.reindex(pd.date_range(ini, df.index.max(), freq=FRECUENCIA, name="Fecha"))
    return salida
