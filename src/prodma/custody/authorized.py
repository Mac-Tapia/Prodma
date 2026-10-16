"""Catálogo autorizado de medidores desde `medidoralimentador.xlsx` (solo lectura).

Identidad exacta: solo se eliminan espacios de borde; guiones y demás caracteres se
conservan. Duplicados exactos con mismos atributos se colapsan y registran todas sus
filas de origen; duplicados con atributos contradictorios bloquean la autorización.
"""
from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pandas as pd

from prodma.excel_schema import normalize_header

SHEET = "alimentador_medidor"
MEDIDOR_HEADER = "medidor"
ALIMENTADOR_HEADER = "codigo alimentador"
NIVEL_HEADER = "nivel tension"
SIGLAS_HEADER = "siglas"
HEADER_SEARCH_ROWS = 20


class AuthorizedCatalogError(ValueError):
    pass


@dataclass(frozen=True)
class AuthorizedCatalog:
    table: pd.DataFrame
    summary: dict


def _find_header(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        normalized = [normalize_header(cell) if cell is not None else "" for cell in row]
        if MEDIDOR_HEADER in normalized:
            return index, {
                "medidor": normalized.index(MEDIDOR_HEADER),
                "alimentador": normalized.index(ALIMENTADOR_HEADER) if ALIMENTADOR_HEADER in normalized else None,
                "nivel": normalized.index(NIVEL_HEADER) if NIVEL_HEADER in normalized else None,
                "siglas": normalized.index(SIGLAS_HEADER) if SIGLAS_HEADER in normalized else None,
            }
    raise AuthorizedCatalogError("No se encontró el encabezado 'Medidor' en las primeras filas")


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def build_authorized_catalog(path: Path | str) -> AuthorizedCatalog:
    path = Path(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in workbook.sheetnames:
            raise AuthorizedCatalogError(f"Falta la hoja '{SHEET}' en {path.name}")
        rows = [tuple(r) for r in workbook[SHEET].iter_rows(values_only=True)]
    finally:
        workbook.close()

    header_index, columns = _find_header(rows)
    records: dict[str, dict] = {}
    data_rows = 0
    for excel_row, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        identifier = _text(row[columns["medidor"]]) if columns["medidor"] < len(row) else None
        if identifier is None:
            continue
        data_rows += 1
        attributes = {
            "codigo_alimentador": _text(row[columns["alimentador"]]) if columns["alimentador"] is not None else None,
            "nivel_tension": _text(row[columns["nivel"]]) if columns["nivel"] is not None else None,
            "siglas": _text(row[columns["siglas"]]) if columns["siglas"] is not None else None,
        }
        if identifier not in records:
            records[identifier] = {"identificador": identifier, **attributes, "filas_origen": [excel_row]}
            continue
        current = records[identifier]
        if any(current[key] != attributes[key] for key in attributes):
            raise AuthorizedCatalogError(
                f"conflicto de atributos para el medidor {identifier!r} "
                f"(filas {current['filas_origen']} y {excel_row})"
            )
        current["filas_origen"].append(excel_row)

    table = pd.DataFrame(sorted(records.values(), key=lambda r: r["identificador"]))
    if table.empty:
        raise AuthorizedCatalogError("El catálogo no contiene medidores")
    summary = {
        "archivo": path.name,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "hoja": SHEET,
        "fila_encabezado": header_index + 1,
        "filas_origen": data_rows,
        "identificadores_unicos": len(table),
        "duplicados_exactos": data_rows - len(table),
    }
    return AuthorizedCatalog(table, summary)


def write_authorized_catalog(catalog: AuthorizedCatalog, out_dir: Path | str) -> dict:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / "authorized_meters.parquet"
    catalog.table.to_parquet(target, index=False)
    reopened = pd.read_parquet(target)
    if len(reopened) != catalog.summary["identificadores_unicos"]:
        raise AuthorizedCatalogError("El parquet publicado no reabre con el número esperado de medidores")
    return {
        "parquet": str(target),
        "sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
        "summary": catalog.summary,
    }
